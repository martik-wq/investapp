"""Export všetkých výkazov firmy do jedného Excelu (.xlsx): výkaz ziskov a strát, súvaha, cash flow,
metriky a valuácia. Hodnoty v miliónoch meny (EPS a pomery v pôvodných jednotkách)."""
import io
import subprocess
import sys
from datetime import date

ORANGE, TINT = "C4570F", "FBEEE2"


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        # po aktualizácii priamo z appky ešte nemusí byť nainštalovaný, skúsime ho doinštalovať
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", "openpyxl"], check=False, timeout=180)
    import openpyxl
    return openpyxl


def _rows(ctx, key):
    """[{'label','vals','kind','bold','indent','header','sub'}] z podrobných alebo zjednodušených výkazov."""
    det = ctx.get("det")
    if det:
        return [{"label": r["label"], "vals": r.get("vals", {}), "kind": r.get("kind", "money"),
                 "bold": bool(r.get("bold")), "indent": r.get("indent", 0) or 0,
                 "header": bool(r.get("header")), "sub": bool(r.get("sub"))} for r in det[key]]
    unit_kind = {"cur": "money", "shares": "shares", "eps": "eps"}
    return [{"label": r["label"], "vals": r["vals"], "kind": unit_kind.get(r["unit"], "money"),
             "bold": bool(r["bold"]), "indent": r["indent"] or 0, "header": False, "sub": False}
            for r in ctx["stm"][key]]


def _scaled(v, kind):
    if v is None:
        return None
    return v if kind == "eps" else v / 1e6


def build_workbook(ctx):
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    comp, built, result = ctx["comp"], ctx["built"], ctx["result"]
    ends = list(reversed(built["display"] if not ctx.get("det") else ctx["det"]["ends"]))
    cur = built["currency"]
    wb = openpyxl.Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=ORANGE)
    sec_fill = PatternFill("solid", fgColor=TINT)
    NUM = {"money": '#,##0.0;[Red]-#,##0.0', "shares": '#,##0.0', "eps": '#,##0.00;[Red]-#,##0.00'}

    # ---- informácie
    ws = wb.active
    ws.title = "Info"
    info = [("Firma", comp["name"]), ("Ticker", comp["ticker"]), ("Mena výkazov", cur),
            ("Obdobie", "ročné" if ctx["mode"] == "annual" else "kvartálne"),
            ("Typ biznisu", result["type_label"]),
            ("Zobrazenie výkazov", "ako firma uvádza" if ctx.get("det") else "zjednodušené (štandardizované)"),
            ("Jednotky", f"mil. {cur} (EPS v {cur}, pomery v %, počet akcií v mil.)"),
            ("Zdroj", "SEC EDGAR (XBRL)"), ("Vytvorené", date.today().strftime("%d.%m.%Y")),
            ("Poznámka", "Údaje sú automaticky mapované z podaní, pri dôležitých rozhodnutiach ich porovnaj "
                         "s originálnym podaním. Nejde o investičné odporúčanie.")]
    for r, (k, v) in enumerate(info, start=1):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80

    def header(sheet, title):
        sheet.cell(1, 1, f"{comp['name']} ({comp['ticker']}) · {title} · mil. {cur}").font = Font(bold=True, size=13)
        sheet.cell(2, 1, "").fill = head_fill
        for i, e in enumerate(ends, start=2):
            c = sheet.cell(2, i, e)
            c.number_format = "dd.mm.yyyy"
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="right")
        sheet.cell(2, 1).font = head_font
        sheet.column_dimensions["A"].width = 58
        for i in range(2, len(ends) + 2):
            sheet.column_dimensions[get_column_letter(i)].width = 14
        sheet.freeze_panes = "B3"

    # ---- výkazy
    for key, title in (("income", "Výkaz ziskov a strát"), ("balance", "Súvaha"), ("cashflow", "Cash flow")):
        sh = wb.create_sheet(title)
        header(sh, title)
        r = 3
        for row in _rows(ctx, key):
            lab = sh.cell(r, 1, row["label"])
            lab.alignment = Alignment(indent=min(int(row["indent"]) + (1 if row["sub"] else 0), 6), wrap_text=False)
            if row["header"]:
                for i in range(1, len(ends) + 2):
                    sh.cell(r, i).fill = sec_fill
                lab.font = Font(bold=True, color=ORANGE)
            else:
                for i, e in enumerate(ends, start=2):
                    v = _scaled(row["vals"].get(e), row["kind"])
                    c = sh.cell(r, i, v)
                    c.number_format = NUM.get(row["kind"], NUM["money"])
                    if row["bold"]:
                        c.font = Font(bold=True)
                    if row["sub"]:
                        c.font = Font(italic=True, color="7A6A5A")
                if row["bold"]:
                    lab.font = Font(bold=True)
                if row["sub"]:
                    lab.font = Font(italic=True, color="7A6A5A")
            r += 1

    # ---- metriky
    sh = wb.create_sheet("Metriky")
    header(sh, f"Metriky ({result['type_label']})")
    mfmt = {"pct": "0.0%", "x": '0.00"×"', "days": "0", "eps": "0.00", "amt": "#,##0.00"}
    r = 3
    for g in result["groups"]:
        c = sh.cell(r, 1, g["title"])
        c.font = Font(bold=True, color=ORANGE)
        for i in range(1, len(ends) + 2):
            sh.cell(r, i).fill = sec_fill
        r += 1
        for row in g["rows"]:
            sh.cell(r, 1, row["label"])
            for i, e in enumerate(ends, start=2):
                v = row["vals"].get(e)
                f = row["fmt"]
                if v is not None and f in ("money", "num"):
                    v = v / 1e6
                cell = sh.cell(r, i, v)
                cell.number_format = mfmt.get(f, '#,##0.0;[Red]-#,##0.0')
            r += 1

    # ---- valuácia
    if ctx.get("val"):
        sh = wb.create_sheet("Valuácia")
        sh.cell(1, 1, f"Valuácia pri cene {ctx['price_used']:.2f} {cur}").font = Font(bold=True, size=13)
        sh.column_dimensions["A"].width = 46
        sh.column_dimensions["B"].width = 18
        for r, (label, v, f) in enumerate(ctx["val"]["items"], start=2):
            sh.cell(r, 1, label)
            if v is not None and f in ("money", "num"):
                v = v / 1e6
            c = sh.cell(r, 2, v)
            c.number_format = mfmt.get(f, '#,##0.0')

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ================================================================== fundamentálna analýza
def _pct_cell(c):
    c.number_format = "0.0%"


def build_fundamental_workbook(ctx):
    """Excel s fundamentálnou analýzou: súhrn, kvalita, modely so živými vzorcami (DCF, P/B, P/E) a konkurencia.
    Vstupy modelov sú v zvýraznených bunkách, po ich zmene sa výsledok v Exceli prepočíta sám."""
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fund, comp = ctx["fund"], ctx["comp"]
    cur = fund["currency"]
    price = fund.get("price")
    wb = openpyxl.Workbook()
    H = Font(bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor=ORANGE)
    SF = PatternFill("solid", fgColor=TINT)
    INP = PatternFill("solid", fgColor="FFE9CC")
    TITLE = Font(bold=True, size=13)
    NOTE = Font(italic=True, color="7A6A5A")

    def head(ws, row, labels, widths=None):
        for i, t in enumerate(labels, start=1):
            c = ws.cell(row, i, t)
            c.font, c.fill = H, HF
            c.alignment = Alignment(horizontal="left" if i == 1 else "right", wrap_text=True)
        if widths:
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(i)].width = w

    # ---- Súhrn
    ws = wb.active
    ws.title = "Súhrn"
    ws["A1"] = f"{comp['name']} ({comp['ticker']}) · fundamentálna analýza"
    ws["A1"].font = TITLE
    ws["A2"] = f"Mena: {cur} · cena na trhu: " + (f"{price:.2f}" if price else "–")
    ws["A3"] = "Rozpätie hodnoty na akciu podľa modelov (pesimistický / základný / optimistický variant)."
    ws["A3"].font = NOTE
    head(ws, 5, ["Model", "Pesimistický", "Základný", "Optimistický", "Základný vs. cena"], [34, 16, 16, 16, 18])
    names = {"dcf": "DCF", "pb": "Férové P/B", "pe": "Cieľové P/E"}
    r = 6
    names.update({"evebitda": "EV / EBITDA", "fcfy": "FCF výnos", "ps": "Cieľové P/S", "ddm": "Dividendy (DDM)", "graham": "Graham (hrubý odhad)"})
    for key in ("dcf", "evebitda", "fcfy", "pe", "ps", "ddm", "pb", "graham"):
        sc = fund[key]["scenarios"]
        if len(sc) != 3 or any(s.get("value") is None for s in sc):
            continue
        ws.cell(r, 1, names[key])
        for i, s in enumerate(sc, start=2):
            ws.cell(r, i, s["value"]).number_format = "#,##0.00"
        if price:
            c = ws.cell(r, 5, sc[1]["value"] / price - 1)
            _pct_cell(c)
        r += 1
    ws.cell(r + 1, 1, "Nejde o investičné odporúčanie. Modely stoja na predpokladoch, ktoré si môžeš upraviť v hárkoch modelov.").font = NOTE

    # ---- Kvalita
    q = fund.get("quality")
    wq = wb.create_sheet("Kvalita")
    wq["A1"] = "Kvalita podniku"
    wq["A1"].font = TITLE
    if not q or q.get("skip"):
        wq["A3"] = "Pri bankách a poisťovniach sa tieto modely nepoužívajú."
    else:
        head(wq, 3, ["Hodnotenie / ukazovateľ", "Hodnota", "Skóre"], [46, 22, 14])
        r = 4
        sk = {2: "silné", 1: "priemerné", 0: "slabé"}
        for g in q["ratings"]:
            c = wq.cell(r, 1, f"{g['name']}: {g['label']}")
            c.font = Font(bold=True, color=ORANGE)
            for i in range(1, 4):
                wq.cell(r, i).fill = SF
            r += 1
            for name, val, sc in g["parts"]:
                wq.cell(r, 1, name).alignment = Alignment(indent=1)
                wq.cell(r, 2, val).alignment = Alignment(horizontal="right")
                wq.cell(r, 3, sk.get(sc, "–")).alignment = Alignment(horizontal="right")
                r += 1
        r += 1
        if q.get("pio"):
            wq.cell(r, 1, f"Piotroski F-score: {q['pio']['score']} z {q['pio']['n']}").font = Font(bold=True)
            r += 1
            for name, ok, why in q["pio"]["items"]:
                wq.cell(r, 1, name).alignment = Alignment(indent=1)
                wq.cell(r, 2, "áno" if ok else ("nie" if ok is not None else "–")).alignment = Alignment(horizontal="right")
                wq.cell(r, 3, why).font = NOTE
                r += 1
            r += 1
        if q.get("alt"):
            wq.cell(r, 1, f"Altman Z'': {q['alt']['z']:.2f} ({q['alt']['zone']} zóna)").font = Font(bold=True)
            r += 1
            for name, v in q["alt"]["parts"]:
                wq.cell(r, 1, name).alignment = Alignment(indent=1)
                wq.cell(r, 2, v).number_format = "0.00"
                r += 1
            r += 1
        if q.get("roic"):
            wq.cell(r, 1, "ROIC").font = Font(bold=True)
            c = wq.cell(r, 2, q["roic"]["roic"])
            _pct_cell(c)
            r += 2
        if q.get("dupont"):
            wq.cell(r, 1, "Rozklad ROE (DuPont)").font = Font(bold=True, color=ORANGE)
            r += 1
            years = q["dupont"]
            wq.cell(r, 1, "").fill = HF
            for i, d in enumerate(years, start=2):
                c = wq.cell(r, i, d["end"].year)
                c.font, c.fill = H, HF
            r += 1
            for lab, key, fmt in (("Čistá marža", "margin", "0.0%"), ("Obrat aktív", "turnover", '0.00"×"'),
                                  ("Finančná páka", "leverage", '0.00"×"'), ("ROE", "roe", "0.0%")):
                wq.cell(r, 1, lab)
                for i, d in enumerate(years, start=2):
                    wq.cell(r, i, d[key]).number_format = fmt
                r += 1

    # ---- DCF so vzorcami
    d = fund["dcf"]
    wd = wb.create_sheet("DCF")
    wd["A1"] = "DCF: diskontovaný voľný cash flow (vstupy v oranžových bunkách, vzorce sa prepočítajú)"
    wd["A1"].font = TITLE
    head(wd, 3, ["Vstup", "Hodnota", "Odporúčané", "Z čoho vychádza"], [44, 16, 16, 90])
    p = d["p"]
    f_by = {f["name"]: f for f in d["fields"]}

    def inp(row, label, val, rec, why, fmt):
        wd.cell(row, 1, label)
        c = wd.cell(row, 2, val)
        c.fill = INP
        c.number_format = fmt
        if rec is not None:
            wd.cell(row, 3, rec).number_format = fmt
        wd.cell(row, 4, why).alignment = Alignment(wrap_text=True)

    def why(name):
        f = f_by[name]
        return " · ".join([x for x in [f.get("why")] + list(f.get("hints") or []) if x])
    inp(4, "Rast tržieb v rokoch 1 až 5", p["g1"], f_by["d_g1"]["rec"] / 100, why("d_g1"), "0.0%")
    inp(5, "FCF marža v roku 5", p["margin5"], f_by["d_m5"]["rec"] / 100, why("d_m5"), "0.0%")
    inp(6, "Diskontná sadzba (WACC)", p["wacc"], f_by["d_wacc"]["rec"] / 100, why("d_wacc"), "0.0%")
    inp(7, "Terminálny rast", p["g_term"], f_by["d_gt"]["rec"] / 100, why("d_gt"), "0.0%")
    inp(8, "Ročné riedenie akcií", p["dilution"], f_by["d_dil"]["rec"] / 100, why("d_dil"), "0.0%")
    inp(9, f"Základ tržieb (mil. {cur})", p["rev0"] / 1e6, f_by["d_rev"]["rec"], why("d_rev"), "#,##0.0")
    inp(10, f"Čistý dlh (mil. {cur})", p["net_debt"] / 1e6, f_by["d_nd"]["rec"], why("d_nd"), "#,##0.0;[Red]-#,##0.0")
    inp(11, "Počet akcií (mil.)", p["shares"] / 1e6, f_by["d_sh"]["rec"], why("d_sh"), "#,##0.00")
    inp(12, "Dnešná FCF marža (východisko)", p["margin0"], None, "Posledných 12 mesiacov (alebo posledný fiškálny rok).", "0.0%")
    inp(13, f"Cena akcie ({cur})", price, None, "Cena na trhu (môžeš prepísať).", "#,##0.00")
    head(wd, 16, ["Rok", "Rast", "Tržby (mil.)", "FCF marža", "FCF (mil.)", "Diskontný faktor", "Dnešná hodnota FCF"])
    for t in range(1, 11):
        r = 16 + t
        wd.cell(r, 1, t)
        wd.cell(r, 2, f"=IF(A{r}<=5,$B$4,$B$4+($B$7-$B$4)*(A{r}-5)/5)").number_format = "0.0%"
        wd.cell(r, 3, f"=$B$9*(1+B{r})" if t == 1 else f"=C{r - 1}*(1+B{r})").number_format = "#,##0.0"
        wd.cell(r, 4, f"=$B$12+($B$5-$B$12)*MIN(A{r},5)/5").number_format = "0.0%"
        wd.cell(r, 5, f"=C{r}*D{r}").number_format = "#,##0.0;[Red]-#,##0.0"
        wd.cell(r, 6, f"=1/(1+$B$6)^A{r}").number_format = "0.0000"
        wd.cell(r, 7, f"=E{r}*F{r}").number_format = "#,##0.0;[Red]-#,##0.0"
    res = [(29, "Súčet dnešných hodnôt FCF (mil.)", "=SUM(G17:G26)", "#,##0.0"),
           (30, "Hodnota po roku 10 (mil.)", "=E26*(1+B7)/(B6-B7)", "#,##0.0"),
           (31, "Dnešná hodnota po roku 10 (mil.)", "=B30*F26", "#,##0.0"),
           (32, "Hodnota firmy (EV, mil.)", "=B29+B31", "#,##0.0"),
           (33, "Hodnota vlastného imania (mil.)", "=B32-B10", "#,##0.0"),
           (34, "Počet akcií po riedení (mil.)", "=B11*(1+B8)^5", "#,##0.00"),
           (35, f"Hodnota na akciu ({cur})", "=B33/B34", "#,##0.00"),
           (36, "Oproti cene", '=IF(B13>0,B35/B13-1,"")', "0.0%")]
    for row, lab, formula, fmt in res:
        wd.cell(row, 1, lab).font = Font(bold=row == 35)
        c = wd.cell(row, 2, formula)
        c.number_format = fmt
        c.font = Font(bold=row == 35)
    wd.cell(38, 1, "Varianty (hodnoty z aplikácie pri tvojich vstupoch)").font = Font(bold=True, color=ORANGE)
    head(wd, 39, ["Variant", "Rast", "FCF marža (rok 5)", "WACC", "Hodnota na akciu"])
    for i, s in enumerate(d["scenarios"]):
        r = 40 + i
        wd.cell(r, 1, s["name"])
        for j, k in enumerate(("g1", "margin5", "wacc"), start=2):
            wd.cell(r, j, s[k]).number_format = "0.0%"
        wd.cell(r, 5, s["value"]).number_format = "#,##0.00"
    if d.get("sens"):
        wd.cell(44, 1, "Citlivosť: hodnota na akciu podľa WACC (riadky) a terminálneho rastu (stĺpce)").font = Font(bold=True, color=ORANGE)
        wd.cell(45, 1, "WACC \\ rast").font = H
        wd.cell(45, 1).fill = HF
        for j, g in enumerate(d["sens"]["gts"], start=2):
            c = wd.cell(45, j, g)
            c.number_format = "0.0%"
            c.font, c.fill = H, HF
        for i, w in enumerate(d["sens"]["waccs"]):
            wd.cell(46 + i, 1, w).number_format = "0.0%"
            for j, v in enumerate(d["sens"]["grid"][i], start=2):
                wd.cell(46 + i, j, v).number_format = "#,##0.00"
    wd.freeze_panes = "A4"

    # ---- P/B
    b = fund["pb"]
    wp = wb.create_sheet("Férové P-B")
    wp["A1"] = "Férové P/B: (ROE − g) / (ke − g) × BVPS (vstupy v oranžových bunkách)"
    wp["A1"].font = TITLE
    head(wp, 3, ["Vstup", "Hodnota", "Odporúčané", "Z čoho vychádza"], [44, 16, 16, 90])
    fb = {f["name"]: f for f in b["fields"]}
    rows = [("Udržateľné ROE", b["p"]["roe"], fb["b_roe"]["rec"] / 100, "b_roe", "0.0%"),
            ("Požadovaná návratnosť (cost of equity)", b["p"]["ke"], fb["b_ke"]["rec"] / 100, "b_ke", "0.0%"),
            ("Dlhodobý rast", b["p"]["g"], fb["b_g"]["rec"] / 100, "b_g", "0.0%"),
            (f"Účtovná hodnota na akciu ({cur})", b["p"]["bvps"], fb["b_bv"]["rec"], "b_bv", "#,##0.00")]
    for i, (lab, val, rec, key, fmt) in enumerate(rows, start=4):
        wp.cell(i, 1, lab)
        c = wp.cell(i, 2, val)
        c.fill, c.number_format = INP, fmt
        if rec is not None:
            wp.cell(i, 3, rec).number_format = fmt
        wp.cell(i, 4, " · ".join([x for x in [fb[key].get("why")] + list(fb[key].get("hints") or []) if x])).alignment = Alignment(wrap_text=True)
    wp.cell(9, 1, "Férové P/B").font = Font(bold=True)
    wp.cell(9, 2, "=(B4-B6)/(B5-B6)").number_format = '0.00"×"'
    wp.cell(10, 1, f"Hodnota na akciu ({cur})").font = Font(bold=True)
    c = wp.cell(10, 2, "=B9*B7")
    c.number_format, c.font = "#,##0.00", Font(bold=True)
    wp.cell(11, 1, "Cena akcie")
    wp.cell(11, 2, price).number_format = "#,##0.00"
    wp.cell(12, 1, "Oproti cene")
    wp.cell(12, 2, '=IF(B11>0,B10/B11-1,"")').number_format = "0.0%"

    # ---- P/E
    e = fund["pe"]
    we = wb.create_sheet("Cieľové P-E")
    we["A1"] = "Cieľové P/E: EPS × (1 + rast) × P/E / (1 + r) (vstupy v oranžových bunkách)"
    we["A1"].font = TITLE
    head(we, 3, ["Vstup", "Hodnota", "Odporúčané", "Z čoho vychádza"], [44, 16, 16, 90])
    fe = {f["name"]: f for f in e["fields"]}
    rows = [(f"EPS za 12 mesiacov ({cur})", e["p"]["eps0"], fe["p_eps"]["rec"], "p_eps", "#,##0.00"),
            ("Rast EPS na budúci rok", e["p"]["g_eps"], fe["p_g"]["rec"] / 100, "p_g", "0.0%"),
            ("Cieľové P/E", e["p"]["target_pe"], fe["p_pe"]["rec"], "p_pe", '0.0"×"'),
            ("Diskontná sadzba", e["p"]["r"], fe["p_r"]["rec"] / 100, "p_r", "0.0%")]
    for i, (lab, val, rec, key, fmt) in enumerate(rows, start=4):
        we.cell(i, 1, lab)
        c = we.cell(i, 2, val)
        c.fill, c.number_format = INP, fmt
        if rec is not None:
            we.cell(i, 3, rec).number_format = fmt
        we.cell(i, 4, " · ".join([x for x in [fe[key].get("why")] + list(fe[key].get("hints") or []) if x])).alignment = Alignment(wrap_text=True)
    we.cell(9, 1, "EPS o rok")
    we.cell(9, 2, "=B4*(1+B5)").number_format = "#,##0.00"
    we.cell(10, 1, "Hodnota o rok")
    we.cell(10, 2, "=B9*B6").number_format = "#,##0.00"
    we.cell(11, 1, f"Hodnota na akciu dnes ({cur})").font = Font(bold=True)
    c = we.cell(11, 2, "=B10/(1+B7)")
    c.number_format, c.font = "#,##0.00", Font(bold=True)
    we.cell(12, 1, "Cena akcie")
    we.cell(12, 2, price).number_format = "#,##0.00"
    we.cell(13, 1, "Oproti cene")
    we.cell(13, 2, '=IF(B12>0,B11/B12-1,"")').number_format = "0.0%"

    # ---- ďalšie modely (hodnoty z aplikácie)
    extra = [m for m in fund.get("models", []) if m["key"] in ("evebitda", "fcfy", "ps", "ddm", "graham")]
    if extra:
        wx = wb.create_sheet("Ďalšie modely")
        wx["A1"] = "Ďalšie oceňovacie modely (vstupy a výsledky z aplikácie, pesimistický / základný / optimistický variant)"
        wx["A1"].font = TITLE
        wx.column_dimensions["A"].width = 44
        for c in "BCD":
            wx.column_dimensions[c].width = 18
        r = 3
        for m in extra:
            wx.cell(r, 1, m["name"]).font = Font(bold=True, color=ORANGE, size=12)
            for i in range(1, 5):
                wx.cell(r, i).fill = SF
            r += 1
            wx.cell(r, 1, "Vstup").font = Font(bold=True)
            wx.cell(r, 2, "Hodnota").font = Font(bold=True)
            wx.cell(r, 3, "Odporúčané").font = Font(bold=True)
            r += 1
            for f in fund[m["key"]]["fields"]:
                wx.cell(r, 1, f"{f['label']} ({f['unit']})")
                wx.cell(r, 2, f["value"]).fill = INP
                if f["rec"] is not None:
                    wx.cell(r, 3, f["rec"])
                r += 1
            r += 1
            wx.cell(r, 1, "Variant").font = Font(bold=True)
            wx.cell(r, 2, f"Hodnota na akciu ({cur})").font = Font(bold=True)
            wx.cell(r, 3, "Oproti cene").font = Font(bold=True)
            r += 1
            for sc in m["scenarios"]:
                wx.cell(r, 1, sc["name"])
                if sc["value"] is not None:
                    wx.cell(r, 2, sc["value"]).number_format = "#,##0.00"
                if sc.get("upside") is not None:
                    _pct_cell(wx.cell(r, 3, sc["upside"]))
                r += 1
            r += 2
        if fund.get("reverse"):
            wx.cell(r, 1, "Reverzný DCF: rast tržieb v rokoch 1 až 5, ktorý je v dnešnej cene").font = Font(bold=True)
            _pct_cell(wx.cell(r, 2, fund["reverse"]["g"]))

    # ---- Konkurencia
    pe_ = fund.get("peers") or {}
    if pe_.get("table"):
        wc = wb.create_sheet("Konkurencia")
        wc["A1"] = "Porovnanie s konkurenciou"
        wc["A1"].font = TITLE
        head(wc, 3, ["Ukazovateľ"] + [c["ticker"] for c in pe_["cols"]], [38] + [16] * len(pe_["cols"]))
        for i, row in enumerate(pe_["table"], start=4):
            wc.cell(i, 1, row["label"])
            for j, v in enumerate(row["vals"], start=2):
                if v is None:
                    continue
                c = wc.cell(i, j, v / 1e6 if row["fmt"] in ("money", "big") else v)
                c.number_format = {"pct": "0.0%", "x": '0.0"×"', "int": "0", "num2": "0.00"}.get(row["fmt"], "#,##0")
                if row["best"] == j - 2:
                    c.fill = SF
                    c.font = Font(bold=True, color=ORANGE)
        wc.cell(4 + len(pe_["table"]) + 1, 1, "Tržby a kapitalizácia v mil. meny výkazov každej firmy. Oranžovo = najlepšia hodnota v riadku.").font = NOTE

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
